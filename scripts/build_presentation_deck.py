#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate institutional-grade Business Plan PPTX for 邻里GEO & 智能增效
Architecture: Decoupled "Text Explanation" + "Dedicated Single Picture Showcase (一图一页, 0 外框, 0 遮挡横线)"
Theme: Brand Purple (#6B21A8, #7E22CE, #9333EA) + Pure White (#FFFFFF) / Light Tint (#FAF5FF)
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# Directories
PROJECTS_DIR = "/Users/a1/代码/GEO"
PICS_DIR = "/Users/a1/Pictures/next 产品图片"
OUTPUT_PPTX = "/Users/a1/Pictures/邻里GEO_智能增效商业计划书.pptx"

# Color Palette (Purple + White strictly)
BG_LIGHT = RGBColor(250, 245, 255)       # #FAF5FF (Soft purple background)
WHITE = RGBColor(255, 255, 255)          # #FFFFFF (Card background)
PURPLE_DARK = RGBColor(88, 28, 135)      # #581C87 (Purple 900, primary titles)
PURPLE_PRIMARY = RGBColor(107, 33, 168)  # #6B21A8 (Purple 800, main brand)
PURPLE_ACCENT = RGBColor(147, 51, 234)   # #9333EA (Purple 600, active/highlight)
PURPLE_BORDER = RGBColor(233, 213, 255)  # #E9D5FF (Card border)
PURPLE_LIGHT_BG = RGBColor(243, 232, 255)# #F3E8FF (Subtle badge background)
TEXT_DARK = RGBColor(15, 23, 42)         # #0F172A (Slate 900)
TEXT_MUTED = RGBColor(71, 85, 105)       # #475569 (Slate 600)

def init_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def set_slide_background(slide, color=BG_LIGHT):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category="NEXTGEO & 小毛驴 AI", subtitle=None):
    """
    Adds a clean, modern header without any horizontal divider lines (0% text blockage).
    """
    # Category tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.28))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_top = 0
    tf_cat.margin_bottom = 0
    tf_cat.margin_left = 0
    tf_cat.margin_right = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = PURPLE_ACCENT

    # Main Title & Subtitle in ONE unified text frame
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.733), Inches(0.75))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.margin_left = 0
    tf.margin_right = 0
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK

    if subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_before = Pt(4)

def add_card(slide, left, top, width, height, fill_color=WHITE, border_color=PURPLE_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
    else:
        card.line.fill.background()
    return card

def convert_to_png_if_needed(path):
    if not os.path.exists(path):
        return None
    if path.lower().endswith('.webp'):
        out_png = f"/tmp/{os.path.basename(path)}.png"
        im = Image.open(path)
        im.convert("RGB").save(out_png)
        return out_png
    return path

def fit_image_in_box(slide, image_path, box_left, box_top, box_width, box_height):
    """
    Fits an image strictly within [box_left, box_top, box_width, box_height],
    maintaining aspect ratio and centering it. Guarantees 0% overflow.
    No outer border or artificial frame is added.
    """
    valid_path = convert_to_png_if_needed(image_path)
    if not valid_path or not os.path.exists(valid_path):
        print(f"Warning: Image not found: {image_path}")
        return None
    try:
        with Image.open(valid_path) as im:
            orig_w, orig_h = im.size

        scale = min(box_width / orig_w, box_height / orig_h)
        disp_w = orig_w * scale
        disp_h = orig_h * scale

        disp_left = box_left + (box_width - disp_w) / 2
        disp_top = box_top + (box_height - disp_h) / 2

        return slide.shapes.add_picture(valid_path, disp_left, disp_top, width=disp_w, height=disp_h)
    except Exception as e:
        print(f"Error placing {image_path}: {e}")
        return None

def build_deck():
    prs = init_presentation()
    blank_layout = prs.slide_layouts[6]

    # Standard positions
    CONTENT_TOP = Inches(1.5)
    CONTENT_HEIGHT = Inches(5.6)

    # =========================================================================
    # SLIDE 1: 封面 (Cover)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, BG_LIGHT)

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PURPLE_PRIMARY
    bar.line.fill.background()

    # Logo
    fit_image_in_box(s1, "/Users/a1/代码/GEO/projects/nextgeo/outputs/assets/logo.jpg", Inches(0.8), Inches(0.8), Inches(1.2), Inches(1.2))

    # Main Card
    add_card(s1, Inches(0.8), Inches(2.2), Inches(11.733), Inches(4.7))

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(11.0), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "2026 战略商业计划书 ｜ 寻找徐州核心资源合伙人"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = PURPLE_ACCENT
    p0.space_after = Pt(12)

    p1 = tf.add_paragraph()
    p1.text = "邻里GEO & 智能增效（小毛驴 AI）"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = PURPLE_DARK
    p1.space_after = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "把大模型重塑为企业的终极信任背调与专属 AI 员工"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = PURPLE_PRIMARY
    p2.space_after = Pt(22)

    bullet_texts = [
        "【邻里GEO】固定 1 万元/年：守护企业数十万投流成本的“终极信任保险”，截流临门一脚",
        "【小毛驴 AI】专属定制 AI 员工：固化老板与销冠思维，不教写 Prompt，员工安装手机 App 即用",
        "【极致省钱】纯轻量定制架构，算力成本自控可本地部署；个人 OPC 仅 20 元/月极简套餐",
        "【工业化交付】AI 自动化全流程辅助 ＋ 徐州/连云港/宿迁京东创业基地高校管培直输入企"
    ]
    for b in bullet_texts:
        pb = tf.add_paragraph()
        pb.text = f"•  {b}"
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(6)

    p_bottom = tf.add_paragraph()
    p_bottom.text = "官方主站：https://www.baicl.cc   ｜   项目主体：徐州璇源网络科技有限公司"
    p_bottom.font.size = Pt(12)
    p_bottom.font.color.rgb = TEXT_MUTED
    p_bottom.space_before = Pt(14)

    # =========================================================================
    # SLIDE 2: 时代剧变：企业营销与 AI 落地两大致命死穴 (痛点分析)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, BG_LIGHT)
    add_header(s2, "时代剧变：企业营销与 AI 落地两大致命死穴", "MACRO SHIFT & INDUSTRY PAIN POINTS", "传统买量遭遇临门一脚截流 ｜ 通用 AI 软件无法融入真实业务")

    card_w = Inches(5.7)
    # Left Card
    add_card(s2, Inches(0.8), CONTENT_TOP, card_w, CONTENT_HEIGHT)
    b1 = s2.shapes.add_textbox(Inches(1.1), CONTENT_TOP + Inches(0.3), card_w - Inches(0.6), CONTENT_HEIGHT - Inches(0.6))
    tf1 = b1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "痛点 1：传统买量在“临门一脚”被 AI 截流"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(14)

    bullets1 = [
        "流量成本十年暴涨：百度竞价、美团外卖、短视频流费水涨船高，获客单价飙升，买量难以为继。",
        "决策链发生不可逆转移：高净值客户与 B2B 采购在付款签约前，必经一步——打开 DeepSeek、豆包或 Kimi 提问：“这家公司到底怎么样？”",
        "AI 查底细当场劝退：若企业未在大模型建立官方权威答案源，AI 会回答“查无此人”甚至推荐竞品软文。",
        "致命失血：企业花了十几万投流买来的意向线索，在付款前最后 10 分钟被大模型生生截胡劝退！"
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # Right Card
    add_card(s2, Inches(6.8), CONTENT_TOP, card_w, CONTENT_HEIGHT)
    b2 = s2.shapes.add_textbox(Inches(7.1), CONTENT_TOP + Inches(0.3), card_w - Inches(0.6), CONTENT_HEIGHT - Inches(0.6))
    tf2 = b2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "痛点 2：通用 AI 软件全成办公室摆设"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(14)

    bullets2 = [
        "泛 AI 软件无法落地：市面上的 AI 纯卖对话框、教员工写 Prompt。员工嫌繁抗拒用，老板看不到 ROI，最后全沦为闲置摆设。",
        "销冠打法无法沉淀：企业销售经验全在销冠个人脑子里，销冠一走客户和打法全被带走；新人培养 3 个月依然不会开单。",
        "缺乏行业专属思考：通用大模型不懂企业行话，输出的内容假大空，根本进不去销售、文案与法务的实际业务流水线。"
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # =========================================================================
    # SLIDE 3: 实操证据：大模型背调实测对比 (大图直接满屏独占展示，0套娃，0外框)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, BG_LIGHT)

    # Minimalist top category label (no redundant slide title, as the image itself has full title)
    tag_box = s3.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.733), Inches(0.25))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "VISUAL PROOF: GEO BENCHMARK 客观实测全景对比"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = PURPLE_ACCENT

    # Image placed directly with 0 outer card, maximum dimensions
    fit_image_in_box(s3, f"{PICS_DIR}/GEO优化前后效果对比.png", Inches(0.8), Inches(0.55), Inches(11.733), Inches(6.6))

    # =========================================================================
    # SLIDE 4: 核心产品一：邻里GEO 品牌答案源 (文字讲解页)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, BG_LIGHT)
    add_header(s4, "邻里GEO：守护企业投流成本的“终极信任保险”", "PRODUCT 1: GENERATIVE ENGINE OPTIMIZATION", "固定 10,000 元/年 ｜ 客户事实源锁死 ｜ 次年极高续费率与超高纯利")

    gw = Inches(5.7)
    gh = Inches(2.65)
    
    geo_features = [
        ("① 结构化事实源编码（普林斯顿标准）", "针对企业官网与核心资产，构建符合 AI 抓取规范的 llms.txt 与 Schema.org 知识实体。大模型爬虫直接提取标准事实，消除传统花哨动画网页无法被 AI 理解的缺陷。"),
        ("② 跨平台权威证据链锁定", "将官方资质、质保承诺、真实标杆案例文档挂载到高权重权威渠道，形成交叉印证的多方证据链，彻底消除大模型幻觉，封杀负面信息与竞品截流。"),
        ("③ 主动沙盘模拟与反向纠偏巡检", "系统定期在 DeepSeek、豆包、Kimi 等主流大模型模拟真实意向客户的刁钻提问。一旦发现 AI 回答不准或推荐同行，系统立即反向定位事实源缺口并进行工程修复。"),
        ("④ 极具吸引力的次年现金流（ARR）", "首年建立全套知识图谱与证据链；次年客户数据与大模型调用关系牢牢沉淀在系统内，迁移成本极高。每年续收 1 万元，边际成本接近于零，成为极高纯利的底层现金流！")
    ]
    
    coords = [
        (Inches(0.8), CONTENT_TOP),
        (Inches(6.8), CONTENT_TOP),
        (Inches(0.8), CONTENT_TOP + gh + Inches(0.2)),
        (Inches(6.8), CONTENT_TOP + gh + Inches(0.2))
    ]

    for i, (title, text) in enumerate(geo_features):
        x, y = coords[i]
        add_card(s4, x, y, gw, gh)
        b = s4.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), gw - Inches(0.5), gh - Inches(0.4))
        tf = b.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PURPLE_DARK
        p.space_after = Pt(8)

        p_body = tf.add_paragraph()
        p_body.text = text
        p_body.font.size = Pt(13)
        p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 5: 实操展示 1：邻里GEO 工业级后台 · 知识库与文章管理 (一图一页，0外框)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, BG_LIGHT)
    add_header(s5, "实战展示：邻里GEO 工业级后台 · 知识库与文章矩阵中枢", "VISUAL PROOF: GEO SYSTEM INTERFACE", "PC端管理总控台：企业结构化事实源与多智能体文章动态编排管理")

    fit_image_in_box(s5, f"{PICS_DIR}/geo1.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 6: 实操展示 2：邻里GEO 工业级后台 · 实时大模型监控雷达 (一图一页，0外框)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, BG_LIGHT)
    add_header(s6, "实战展示：邻里GEO 工业级后台 · 实时大模型监控雷达", "VISUAL PROOF: GEO SYSTEM INTERFACE", "实时监测 DeepSeek / 豆包 / Kimi / 文心等大模型对企业的可见度与认知事实")

    fit_image_in_box(s6, f"{PICS_DIR}/GEO 2.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 7: 核心产品二：企业专属 AI 员工（小毛驴 AI）(文字讲解页)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, BG_LIGHT)
    add_header(s7, "小毛驴 AI：为企业定制“永不离职、能自我升级”的 AI 员工", "PRODUCT 2: BESPOKE AI EMPLOYEE", "不教员工写 Prompt ｜ 员工手机装 APP 即可使用 ｜ 萃取老板与销冠脑力")

    # Left: Core Logic
    add_card(s7, Inches(0.8), CONTENT_TOP, Inches(4.3), CONTENT_HEIGHT)
    b_left = s7.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), Inches(3.9), CONTENT_HEIGHT - Inches(0.6))
    tf_l = b_left.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "三位一体运作闭环"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    steps = [
        ("萃取老板与销冠思维", "把金牌销售的成交流程与老板的商业直觉萃取固化为专属智能体。"),
        ("越用越聪明·自主升级", "AI 员工会在实际业务对话与客户反馈中持续迭代，自我沉淀经验。"),
        ("轻量客户端·极速上手", "员工无需理解复杂技术，安装手机 APP 点击专属模块即可直接作业。")
    ]
    for stitle, stext in steps:
        p_st = tf_l.add_paragraph()
        p_st.text = f"▶ {stitle}"
        p_st.font.size = Pt(14)
        p_st.font.bold = True
        p_st.font.color.rgb = PURPLE_PRIMARY
        p_st.space_after = Pt(4)

        p_sb = tf_l.add_paragraph()
        p_sb.text = stext
        p_sb.font.size = Pt(12)
        p_sb.font.color.rgb = TEXT_DARK
        p_sb.space_after = Pt(12)

    # Right: 4 Scenarios
    add_card(s7, Inches(5.4), CONTENT_TOP, Inches(7.133), CONTENT_HEIGHT)
    b_right = s7.shapes.add_textbox(Inches(5.7), CONTENT_TOP + Inches(0.3), Inches(6.533), CONTENT_HEIGHT - Inches(0.6))
    tf_r = b_right.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "深度定制四大核心业务场景"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    scenarios = [
        ("1. 公域全自动获客", "批量生成贴合抖音、小红书、视频号算法的高转化原创脚本与图文，吸引精准意向客户。"),
        ("2. 私域高情商朋友圈文案", "告别生硬发广告。根据企业定位与客户画像，一键输出真实、生活化、高粘性的私域内容。"),
        ("3. 智能合同审查与穿透尽调", "秒级排查合同风险条款、违约陷阱，一键穿透合作方企业资质、司法诉讼与失信记录。"),
        ("4. 销冠经验固化与员工培训", "将销冠应对话术与产品知识库打包成陪练助手，新人上岗即可达到 80 分专业水准。")
    ]
    for stitle, stext in scenarios:
        p_st = tf_r.add_paragraph()
        p_st.text = stitle
        p_st.font.size = Pt(14)
        p_st.font.bold = True
        p_st.font.color.rgb = PURPLE_PRIMARY
        p_st.space_after = Pt(3)

        p_sb = tf_r.add_paragraph()
        p_sb.text = stext
        p_sb.font.size = Pt(12)
        p_sb.font.color.rgb = TEXT_DARK
        p_sb.space_after = Pt(10)

    # =========================================================================
    # SLIDE 8: 实操展示：小毛驴 AI 客户端真机交互与功能专享 (大图独立展示，0外框)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, BG_LIGHT)
    add_header(s8, "实战展示：小毛驴 AI 客户端真机交互与功能专享", "VISUAL PROOF: MOBILE APP WORKFLOW", "不同企业专享定制功能 ｜ 冗余功能自动隐藏 ｜ 手机 App 随时随地随身使用")

    phone_w = Inches(3.4)
    phone_h = Inches(5.5)
    y_pos = CONTENT_TOP + Inches(0.05)
    fit_image_in_box(s8, f"{PICS_DIR}/小毛驴 AI 1.jpg", Inches(1.1), y_pos, phone_w, phone_h)
    fit_image_in_box(s8, f"{PICS_DIR}/小毛驴 ai 2.jpg", Inches(4.966), y_pos, phone_w, phone_h)
    fit_image_in_box(s8, f"{PICS_DIR}/小毛驴 AI 3.jpg", Inches(8.833), y_pos, phone_w, phone_h)

    # =========================================================================
    # SLIDE 9: 核心能力三：智能合同审查与穿透尽调 (文字讲解页)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, BG_LIGHT)
    add_header(s9, "智能合同审查与穿透尽调：企业法务风控安全底座", "PRODUCT 3: LEGAL & DUE DILIGENCE", "中小企业以极低成本拥有“随身法务总监” ｜ 守住企业经营与合同底线")

    col_w = Inches(3.7)
    gap = Inches(0.316)

    # Col 1: 痛点
    add_card(s9, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_l1 = s9.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_l1 = b_l1.text_frame
    tf_l1.word_wrap = True
    p = tf_l1.paragraphs[0]
    p.text = "中小微企业法务困局"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_l1.add_paragraph()
    p_body.text = "• 法律顾问贵：年费 3~5 万起步，响应迟缓，小合同不愿看。\n\n• 踩坑成本极高：合同违约金条款漏洞、免责陷阱频发，一次纠纷损失数万元。\n\n• 交易对象难辨真伪：签约前对合作方底细一无所知，被皮包公司或老赖截流欺骗。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # Col 2: 审查与排雷
    add_card(s9, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_l2 = s9.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_l2 = b_l2.text_frame
    tf_l2.word_wrap = True
    p = tf_l2.paragraphs[0]
    p.text = "合同秒级排雷审查"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_l2.add_paragraph()
    p_body.text = "• 秒级全文深度扫描：支持拍照、PDF、Word 直接上传，AI 自动定位隐蔽风险条款。\n\n• 智能出具修改建议：严格对齐《民法典》及相关行业标准，直接生成可签约的合规版本。\n\n• 关键条款高亮标注：付款账期、违约倍数、不可抗力等核心商业利益全面布防。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # Col 3: 穿透尽调
    add_card(s9, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_l3 = s9.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_l3 = b_l3.text_frame
    tf_l3.word_wrap = True
    p = tf_l3.paragraphs[0]
    p.text = "交易主体穿透尽调"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_l3.add_paragraph()
    p_body.text = "• 股东与实际控制人穿透：多层股权穿透、关联企业排查，锁定最终幕后实控人。\n\n• 司法涉诉与失信排查：实时调取历史裁判文书、限制高消费、被执行人等风险记录。\n\n• 经营异常风险雷达：行政处罚、税收违法、工商异常经营状态一键全景呈现。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 10: 实操展示 1：智能法务 · 合同秒级风险排查与条款审查 (一图一页，0外框)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, BG_LIGHT)
    add_header(s10, "实战展示：智能法务 · 合同秒级风险排查与条款审查报告", "VISUAL PROOF: LEGAL RISK AUDIT", "自动标注文档风险条款、违约金漏洞与免责陷阱，并出具合规修改建议")

    fit_image_in_box(s10, f"{PICS_DIR}/小毛驴 ai 法务和同审核.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 11: 实操展示 2：智能法务 · 目标企业穿透尽调与风险雷达 (一图一页，0外框)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, BG_LIGHT)
    add_header(s11, "实战展示：智能法务 · 目标企业穿透尽调与风险全景雷达", "VISUAL PROOF: ENTERPRISE DUE DILIGENCE", "多层股权穿透、司法诉讼/失信记录穿透与经营合规深度评估报告")

    fit_image_in_box(s11, f"{PICS_DIR}/小毛驴 ai 法务.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 12: 赛道定力与架构壁垒：做极致省钱的轻量定制 vs 通用大厂 (战略对比)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, BG_LIGHT)
    add_header(s12, "赛道定力与架构壁垒：做极致省钱的轻量定制 vs 通用大厂", "COMPETITIVE ADVANTAGE & ARCHITECTURE", "避开通用 AI 算力军备竞赛 ｜ 算力成本完全可控 ｜ 个人 OPC 仅 20 元/月极具杀伤力")

    col_w = Inches(3.7)
    gap = Inches(0.316)

    # Card 1: 为什么不用大厂
    add_card(s12, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_c1 = s12.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_c1 = b_c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "为什么不用通用大厂软件？"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(10)

    p_body = tf_c1.add_paragraph()
    p_body.text = "• 通用软件 Token 成本不可控：像腾讯 WorkBuddy 等大厂工具面向通用泛场景，未来 Token 消耗与计费必然水涨船高。\n\n• 通用功能冗余且复杂：大厂软件给所有人用同一套复杂界面，菜单上百个，员工看一眼就放弃。\n\n• 我们做纯定制：每个客户进入后，只保留其业务最需要的功能，其他多余功能全部消失！"
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = TEXT_DARK

    # Card 2: 算力自控与本地化
    add_card(s12, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_c2 = s12.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_c2 = b_c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "算力成本自控与私有化部署"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(10)

    p_body = tf_c2.add_paragraph()
    p_body.text = "• 不切通用赛道：通用 AI 军备竞赛消耗极其庞大，而我们只切中小企业最核心的 3~4 个高频商业问题。\n\n• 当前 AI 智商已完全溢出：业务场景无需万亿级通用模型，小模型/端侧模型足矣。\n\n• 极简本地私有化：满足企业“数据不上云、本地专属微调、私有化部署”的高安全需求，极其安全且边际成本趋零！"
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = TEXT_DARK

    # Card 3: 杀手级定价策略
    add_card(s12, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_c3 = s12.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_c3 = b_c3.text_frame
    tf_c3.word_wrap = True
    p = tf_c3.paragraphs[0]
    p.text = "杀手级超低价与高客户粘性"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(10)

    p_body = tf_c3.add_paragraph()
    p_body.text = "• 个人 OPC 创业套餐：仅需 20 元/月！即可满足个体一人公司日常图文、文案、获客全套工作，用量充足直观，降维打击市场。\n\n• 企业版订阅仅 100~200 元/月：大幅低于一名初级文员的薪酬，续费阻力极低。\n\n• 越做客户越稳定：客户用大厂产品随时可换，但在我们这里的定制功能沉淀了其独家语料和打法，客户终身绑定！"
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 13: 微信群即入口与多租户安全隔离 (交互创新)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, BG_LIGHT)
    add_header(s13, "微信群即入口与多租户安全隔离：零门槛沉浸交互", "ARCHITECTURE & SECURITY", "无需改变员工使用习惯 ｜ 独立专属群协同 ｜ 企业核心数据物理隔离")

    add_card(s13, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_w1 = s13.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_w1 = b_w1.text_frame
    tf_w1.word_wrap = True
    p = tf_w1.paragraphs[0]
    p.text = "微信群即工作台"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_w1.add_paragraph()
    p_body.text = "• 零学习成本：员工和老板不需要下载额外办公套件，直接在熟悉的微信群内 @ AI 员工发起任务。\n\n• 自然语言派单：“帮我写一条今晚 8 点美容院拓客的朋友圈”、“排查这份供货合同违约责任”，秒级交付。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s13, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_w2 = s13.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_w2 = b_w2.text_frame
    tf_w2.word_wrap = True
    p = tf_w2.paragraphs[0]
    p.text = "群内降噪与指令沉淀"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_w2.add_paragraph()
    p_body.text = "• 噪音精准过滤：系统自动过滤微信群内闲聊灌水，仅针对 @ 机器人 或触发特定指令的业务诉求做出结构化响应。\n\n• 过程成果沉淀：文案生成记录、合同排雷报告、拓客数据自动归档同步至 PC 后台，老板随时复盘。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s13, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_w3 = s13.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_w3 = b_w3.text_frame
    tf_w3.word_wrap = True
    p = tf_w3.paragraphs[0]
    p.text = "多租户物理强隔离"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_w3.add_paragraph()
    p_body.text = "• 独立房间架构：每个企业拥有唯一的安全隔离房间与加密通道，确保竞品之间绝不串群、语料绝不混淆。\n\n• 商业秘密不出圈：企业的核心价格体系、客户名单与销冠话术仅在私有容器内运转，物理级数据风控。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 14: 实操展示：微信社群中枢与独立多租户交付房间 (大图独立展示，0外框)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, BG_LIGHT)
    add_header(s14, "实战展示：微信社群中枢与独立多租户交付房间", "VISUAL PROOF: WECHAT COMMUNITY HUB", "企业专属社群中枢 ｜ 独立多租户物理隔离 ｜ 手机端自然交互无缝接入")

    fit_image_in_box(s14, f"{PICS_DIR}/社区 1.jpg", Inches(1.1), y_pos, phone_w, phone_h)
    fit_image_in_box(s14, f"{PICS_DIR}/社区 2.jpg", Inches(4.966), y_pos, phone_w, phone_h)
    fit_image_in_box(s14, f"{PICS_DIR}/社区 3.jpg", Inches(8.833), y_pos, phone_w, phone_h)

    # =========================================================================
    # SLIDE 15: 交付网络：全流程 AI 提效 ＋ 高校管培直输入企 (交付保障)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15, BG_LIGHT)
    add_header(s15, "工业化交付体系：全流程 AI 提效 + 徐连宿高校管培入企", "DELIVERY SYSTEM & SCALE EFFICIENCY", "标准化人机协同流水线 ｜ 边际履约成本极低 ｜ 京东创业基地持续输送实战人才")

    add_card(s15, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_d1 = s15.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_d1 = b_d1.text_frame
    tf_d1.word_wrap = True
    p = tf_d1.paragraphs[0]
    p.text = "AI 自动化生产流水线"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_d1.add_paragraph()
    p_body.text = "• 研发人效提升 10 倍：代码脚手架、Schema.org 数据对齐、知识库提纯全部由 AI 自动化脚本秒级完成。\n\n• 极低边际交付成本：传统软件公司定制一个系统需 3~5 人月，我们仅需 1 名工程师 2 天即可完成深度定制并上线。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s15, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_d2 = s15.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_d2 = b_d2.text_frame
    tf_d2.word_wrap = True
    p = tf_d2.paragraphs[0]
    p.text = "徐连宿高校管培基地"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_d2.add_paragraph()
    p_body.text = "• 深厚产教融合积淀：曾统筹负责淘宝联盟在徐州、连云港、宿迁几乎所有大学的校企合作基地，后全面升级更名为京东创业基地。\n\n• 源源不断的实操生力军：数千名在校大学生直接接受现代电商与 AI 工具实操培训，为本地企业输送极高性价比管培生。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s15, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_d3 = s15.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_d3 = b_d3.text_frame
    tf_d3.word_wrap = True
    p = tf_d3.paragraphs[0]
    p.text = "管培生入企驻场辅导"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_d3.add_paragraph()
    p_body.text = "• 手把手带教客户员工：经我们培训合格的高校优秀学生作为交付专员，入驻客户现场帮带员工使用 APP，打消老板落地顾虑。\n\n• 双向共赢：企业获得高素质年轻人操作新工具，大学生获得优质实习就业机会，形成徐州本地坚不可摧的交付壁垒。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 16: 实操展示 1：产教融合 · 徐州/连云港/宿迁高校合作网络 (一图一页，0外框)
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_background(s16, BG_LIGHT)
    add_header(s16, "实战展示：产教融合 · 徐州/连云港/宿迁高校合作网络", "VISUAL PROOF: UNIVERSITY TALENT PIPELINE", "阿里校企合作基地发展沉淀，现已全面升级更名为京东创业基地")

    fit_image_in_box(s16, f"{PICS_DIR}/产教融合创业基地.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 17: 实操展示 2：校企实训 · 产教融合与电商实操现场 (一图一页，0外框)
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    set_slide_background(s17, BG_LIGHT)
    add_header(s17, "实战展示：校企实训 · 产教融合与电商实操签约现场", "VISUAL PROOF: UNIVERSITY TALENT PIPELINE", "标准化管培生直输企业，构建徐州本地极低边际成本的工业化履约流水线")

    fit_image_in_box(s17, f"{PICS_DIR}/校企合作 1.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 18: 实战打样 1：全国一线美业连锁标杆 (业务讲解)
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    set_slide_background(s18, BG_LIGHT)
    add_header(s18, "实战打样 1：全国一线美业连锁标杆", "CASE STUDY: NATIONAL BEAUTY BRANDS", "苗医生、豆域、克丽缇娜等品牌数字化增效 ｜ 支付宝大型营销活动赋能")

    add_card(s18, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_b1 = s18.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_b1 = b_b1.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "一线标杆美业品牌赋能"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_b1.add_paragraph()
    p_body.text = "• 深度服务头部连锁：长期为苗医生全国连锁门店、豆域高端万平美业综合体、克丽缇娜加盟体系提供数字化营销赋能。\n\n• 验证商业模型：美业高度依赖私域复购与到店转化，对文案和公私域获客要求极高，AI 员工落地效果显著。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s18, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_b2 = s18.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_b2 = b_b2.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "定制化 AI 员工核心场景"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_b2.add_paragraph()
    p_body.text = "• 私域朋友圈批量产出：自动生成符合高端调性的朋友圈文案与项目打卡图，店员一键复制。\n\n• 销冠拓客话术固化：美容顾问通过手机 APP 获取最新项目沟通脚本，客单价与到店率提升超 40%。\n\n• 客诉预警排查：敏感辞令自动纠错，杜绝虚假宣传合规风险。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s18, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_b3 = s18.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_b3 = b_b3.text_frame
    tf_b3.word_wrap = True
    p = tf_b3.paragraphs[0]
    p.text = "支付宝大型营销活动赋能"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_b3.add_paragraph()
    p_body.text = "• 公私域深度联动：结合支付宝官方营销节点与本地生活优惠券，实现从线上领券、到店核销到私域沉淀全链路数字化。\n\n• 极高转介绍率：在徐州本地美业圈层形成示范效应，多个同业品牌主动慕名寻求定制合作。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 19: 实战展示：一线美业连锁数字化增效打样 (一图一页，0外框)
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    set_slide_background(s19, BG_LIGHT)
    add_header(s19, "实战展示：一线美业连锁数字化增效打样", "VISUAL PROOF: BEAUTY RETAIL CASE", "苗医生、豆域、克丽缇娜等品牌数字化增效与支付宝大型营销活动赋能")

    fit_image_in_box(s19, f"{PICS_DIR}/美业合作与支付宝活动作为标杆.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 20: 实操展示：实体农业龙头 · 大恒农业数字化增效 (一图一页，0外框)
    # =========================================================================
    s20 = prs.slides.add_slide(blank_layout)
    set_slide_background(s20, BG_LIGHT)
    add_header(s20, "实战打样 2：实体农业龙头 · 大恒农业数字化增效", "CASE STUDY: MODERN AGRICULTURE", "年会员费超 7000 万元的现代农业龙头企业增效与私域营销体系支撑")

    fit_image_in_box(s20, f"{PICS_DIR}/大恒农业一年7000 万会员费.jpg", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 21: 实操展示：现代竞技体育 · 白鲨格斗区域第一拳馆 (一图一页，0外框)
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    set_slide_background(s21, BG_LIGHT)
    add_header(s21, "实战打样 3：现代竞技体育 · 白鲨格斗区域第一拳馆", "CASE STUDY: REGIONAL FIGHT CLUB", "矩阵短视频低成本引流与公私域联动，成就连云港区域第一拳馆标杆")

    fit_image_in_box(s21, f"{PICS_DIR}/白鲨格斗简单的视成就连云港第一拳馆.jpg", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 22: 核心团队：黄金铁三角架构 (团队介绍)
    # =========================================================================
    s22 = prs.slides.add_slide(blank_layout)
    set_slide_background(s22, BG_LIGHT)
    add_header(s22, "核心团队：黄金铁三角架构", "CORE TEAM ARCHITECTURE", "十年大厂技术架构师 ＋ 本地头部流量操盘手 ＋ 高校产教管培导师")

    add_card(s22, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_t1 = s22.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_t1 = b_t1.text_frame
    tf_t1.word_wrap = True
    p = tf_t1.paragraphs[0]
    p.text = "技术与架构中枢"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_t1.add_paragraph()
    p_body.text = "• 创始人 / 技术总监\n\n• 十余年全栈研发与大模型架构经验，曾主导多款百万级用户互联网产品底层架构设计。\n\n• 专精 GEO 普林斯顿算法标准、企业轻量级模型微调与多租户微信交互安全中枢开发。\n\n• 全权主导本系统底层代码自研与知识产权沉淀。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s22, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_t2 = s22.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_t2 = b_t2.text_frame
    tf_t2.word_wrap = True
    p = tf_t2.paragraphs[0]
    p.text = "商业与流量操盘"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_t2.add_paragraph()
    p_body.text = "• 联合创始人 / 商业化总监\n\n• 资深互联网流量操盘手，曾成功操盘惠发现 APP 平台，拥有百万级本地精准消费客群运营经验。\n\n• 深度掌握本地实体商业获客密码与微信私域运营策略，主导美业头部品牌及实体农业数字化打样。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s22, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_t3 = s22.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_t3 = b_t3.text_frame
    tf_t3.word_wrap = True
    p = tf_t3.paragraphs[0]
    p.text = "高校产教与交付中枢"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_t3.add_paragraph()
    p_body.text = "• 交付合伙人 / 产教融合总监\n\n• 原淘宝联盟徐连宿高校电商合作基地、京东创业基地实操负责人，深耕高校产教融合十余年。\n\n• 掌控本地充沛的高校实操人才池，负责将高校管培生源源不断定向输送入企，提供零阻力上门交付与售后培训。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 23: 实操展示：团队资质背书与行业影响力 (一图一页，0外框)
    # =========================================================================
    s23 = prs.slides.add_slide(blank_layout)
    set_slide_background(s23, BG_LIGHT)
    add_header(s23, "团队资质背书：淮海经济区大会与实操案例", "VISUAL PROOF: REGIONAL CREDIBILITY", "深耕徐州及淮海经济区多年，具备顶级行业峰会演讲与大型实战交付履历")

    fit_image_in_box(s23, f"{PICS_DIR}/团队案例淮海经济区大会及可兼经历.png", Inches(0.8), CONTENT_TOP, Inches(11.733), CONTENT_HEIGHT)

    # =========================================================================
    # SLIDE 24: 战略根据地与未来格局：徐州 300~500 家与“本地小米生态” (未来格局)
    # =========================================================================
    s24 = prs.slides.add_slide(blank_layout)
    set_slide_background(s24, BG_LIGHT)
    add_header(s24, "战略根据地与未来格局：徐州 300~500 家企业底座与“本地小米生态”", "FUTURE VISION: LOCAL XIAOMI ECOSYSTEM", "不求全国虚胖铺摊子 ｜ 扎透徐州核心企业圈 ｜ 从工具软件跃迁为本地高净值商业盟友")

    add_card(s24, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_f1 = s24.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_f1 = b_f1.text_frame
    tf_f1.word_wrap = True
    p = tf_f1.paragraphs[0]
    p.text = "阶段一：扎透徐州根据地"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_f1.add_paragraph()
    p_body.text = "• 聚焦目标：用 12~18 个月，深度服务徐州 300~500 家各行各业头部中小企业与高净值老板。\n\n• 形成极高口碑护城河：每个行业仅打造 1~2 个标杆客户，先入为主占领大模型权威推荐位，形成区域同业绝对壁垒。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s24, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_f2 = s24.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_f2 = b_f2.text_frame
    tf_f2.word_wrap = True
    p = tf_f2.paragraphs[0]
    p.text = "阶段二：高纯利飞轮运转"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_f2.add_paragraph()
    p_body.text = "• 底层现金流极稳：300 家客户 × 每年 1 万元基础年费 ＝ 300 万元稳定 ARR 纯利现金流。\n\n• 超高复购续约：因企业核心数字资产与模型调用深度绑定，客户第二年主动续约率预计超 85%，边际交付成本几乎为零！"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s24, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_f3 = s24.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_f3 = b_f3.text_frame
    tf_f3.word_wrap = True
    p = tf_f3.paragraphs[0]
    p.text = "阶段三：本地“小米生态链”"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_f3.add_paragraph()
    p_body.text = "• 终局格局绝非纯卖软件：这 300~500 家企业的老板，是徐州本地最具消费力与商业号召力的高净值人群。\n\n• 商业赋能分发生态：未来向该联盟企业严选输出供应链金融、高端企业内训、联名私域集采，打造本地最强商业闭环！"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 25: 合作共赢与合伙人权益：开放 10% 核心股权 (合作方案)
    # =========================================================================
    s25 = prs.slides.add_slide(blank_layout)
    set_slide_background(s25, BG_LIGHT)
    add_header(s25, "合作共赢与合伙人权益：开放 10% 核心股权", "STRATEGIC PARTNERSHIP & EQUITY OFFER", "寻找深度互信的核心资源合伙人 ｜ 现金流纯分红 ｜ 严格有限责任与债务绝对隔离")

    add_card(s25, Inches(0.8), CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_p1 = s25.shapes.add_textbox(Inches(1.0), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_p1 = b_p1.text_frame
    tf_p1.word_wrap = True
    p = tf_p1.paragraphs[0]
    p.text = "为什么开放 10% 股权？"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_p1.add_paragraph()
    p_body.text = "• 我们不缺纯财务投资，缺的是【本土核心资源操盘手】。\n\n• 期望合伙人能够为项目在徐州本地政府园区、商会圈层、大型集团客户对接提供强背书。\n\n• 双方各展所长：我们负责死磕技术研发与极致交付，合伙人负责本地商业圈层高层撬动与资源护航。"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s25, Inches(0.8) + col_w + gap, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_p2 = s25.shapes.add_textbox(Inches(0.8) + col_w + gap + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_p2 = b_p2.text_frame
    tf_p2.word_wrap = True
    p = tf_p2.paragraphs[0]
    p.text = "出资形式与纯分红权益"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_p2.add_paragraph()
    p_body.text = "• 10% 核心股权出资形式：诚意现金入股 ＋ 核心政商/圈层资源入股（具体入资金额双方线下商定）。\n\n• 极快分红周期：轻资产运作，无大厂巨额硬件折旧包袱。首年实现规模化收费后，按季度/半年度享受真实净利润 10% 现金分红！"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    add_card(s25, Inches(0.8) + (col_w + gap)*2, CONTENT_TOP, col_w, CONTENT_HEIGHT)
    b_p3 = s25.shapes.add_textbox(Inches(0.8) + (col_w + gap)*2 + Inches(0.2), CONTENT_TOP + Inches(0.3), col_w - Inches(0.4), CONTENT_HEIGHT - Inches(0.6))
    tf_p3 = b_p3.text_frame
    tf_p3.word_wrap = True
    p = tf_p3.paragraphs[0]
    p.text = "有限责任·负债绝对隔离"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.space_after = Pt(12)

    p_body = tf_p3.add_paragraph()
    p_body.text = "• 严格有限责任保护：基于《公司法》有限责任公司架构，合伙人仅以其实缴出资额为限承担有限责任。\n\n• 债务终身绝对隔离：签署严密合伙协议，公司一切经营性借款、银行信贷或潜在历史债务由创始团队全权承担，合伙人 100% 免责！"
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 26: 封底页：携手共建徐州企业智能化基础设施 (Closing)
    # =========================================================================
    s26 = prs.slides.add_slide(blank_layout)
    set_slide_background(s26, BG_LIGHT)

    bar26 = s26.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    bar26.fill.solid()
    bar26.fill.fore_color.rgb = PURPLE_PRIMARY
    bar26.line.fill.background()

    add_card(s26, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    b_close = s26.shapes.add_textbox(Inches(2.0), Inches(2.0), Inches(9.333), Inches(3.5))
    tf_c = b_close.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "携手共建徐州企业智能化增长新引擎"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = PURPLE_DARK
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(16)

    p_sub = tf_c.add_paragraph()
    p_sub.text = "邻里GEO ＋ 小毛驴 AI · 让 AI 真正成为企业的现金流保障"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = PURPLE_PRIMARY
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_after = Pt(30)

    p_info = tf_c.add_paragraph()
    p_info.text = "官方主站：https://www.baicl.cc\n项目主体：徐州璇源网络科技有限公司\n合作接洽：项目创始人团队 ｜ 欢迎预约线下实操演示与系统对撞"
    p_info.font.size = Pt(14)
    p_info.font.color.rgb = TEXT_MUTED
    p_info.alignment = PP_ALIGN.CENTER

    # Save
    prs.save(OUTPUT_PPTX)
    print(f"Presentation successfully saved to: {OUTPUT_PPTX} (Total slides: {len(prs.slides)})")

if __name__ == "__main__":
    build_deck()
